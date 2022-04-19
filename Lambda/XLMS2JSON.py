import json
from collections import OrderedDict
from itertools import islice
from openpyxl import load_workbook

# Open the workbook and select a worksheet
wb = load_workbook('/content/data.xlsm')    ##file name here
sheet = wb['Skill Profiles']

# List to hold dictionaries
skill_list = []

# Iterate through each row in worksheet and fetch values into dict
for row in islice(sheet.values, 11, sheet.max_row):
    skill = OrderedDict()
    skill['Skill-Profiles'] = row[4]
    skill['Skill-ID'] = row[5]
    skill['Skill-Category'] = row[6]
    skill['Skillname EN'] = row[9]
    skill['Skill Description EN'] = row[10]
    Levels = OrderedDict()
    Levels['Junior'] = row[11]
    Levels['Specialist'] = row[12]
    Levels['Expert'] = row[13]
    Levels['Senior'] = row[14]
    Levels['Principal'] = row[15]
    skill['Levels']=Levels
    skill_list.append(skill)

# Serialize the list of dicts to JSON
j = json.dumps(cars_list)

#Write to file
with open('data.json', 'w') as f:
    f.write(j)
